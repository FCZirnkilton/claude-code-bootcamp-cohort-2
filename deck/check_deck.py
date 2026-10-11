#!/usr/bin/env python3
"""Structural checks for the presenter deck (run in CI): numbering, notes, videos, prompts, overflow."""

import re
import sys
from pathlib import Path

from pptx import Presentation

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "tools"))
sys.path.insert(0, str(ROOT / "deck"))
from build_deck import wrap_prompt  # noqa: E402
from course_data import PARTS  # noqa: E402

deck = Presentation(ROOT / "deck" / "claude-code-bootcamp-cohort-2.pptx")
total, problems, videos, slide_text = len(deck.slides), [], 0, []
for i, slide in enumerate(deck.slides, start=1):
    texts = [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame]
    slide_text.append("\n".join(texts))
    counter = [t.strip() for t in texts if re.fullmatch(r"\d{3} / \d+", t.strip())]
    if counter and counter[0] != f"{i:03d} / {total}":
        problems.append(f"slide {i}: counter {counter[0]}")
    if not slide.notes_slide.notes_text_frame.text.strip():
        problems.append(f"slide {i}: no speaker notes")
    if any("PENDING" in t or "<next part" in t for t in texts):
        problems.append(f"slide {i}: placeholder text")
    if any(r.reltype.endswith("/video") for r in slide.part.rels.values()):
        videos += 1
    for shape in slide.shapes:
        if shape.top is not None and shape.top >= 0 and shape.top + shape.height > deck.slide_height + 9525 * 10:
            problems.append(f"slide {i}: shape extends below the slide")
        if shape.has_table and shape.top / 914400 + len(shape.table.rows) * 0.36 > 7.4:
            problems.append(f"slide {i}: table with {len(shape.table.rows)} rows overflows")
if videos != len(PARTS):
    problems.append(f"{videos} video slides, expected {len(PARTS)}")
joined = "\n".join(slide_text)
for part in PARTS:
    for _, prompt in part["prompts"]:
        if wrap_prompt(prompt) not in joined:
            problems.append(f"part {part['n']:02d}: prompt on slides differs from tools/course_data.py (rebuild the deck)")
print(f"{total} slides, {videos} videos")
print("\n".join(problems) or "deck OK")
sys.exit(1 if problems else 0)
