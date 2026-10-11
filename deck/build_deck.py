#!/usr/bin/env python3
"""Build deck/claude-code-bootcamp-cohort-2.pptx from deck/content.py, tools/course_data.py and videos/.

Requires python-pptx (pip install python-pptx) and ffmpeg (for video poster frames).
"""

from __future__ import annotations

import subprocess
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "tools"))
sys.path.insert(0, str(ROOT / "deck"))
from content import BREAK_NOTES, CONCEPTS, FINAL, INTRO  # noqa: E402
from course_data import PARTS, schedule  # noqa: E402
from run_facts import facts  # noqa: E402

OUT = ROOT / "deck" / "claude-code-bootcamp-cohort-2.pptx"
POSTERS = ROOT / "deck" / ".posters"
C = lambda h: RGBColor.from_string(h.lstrip("#"))  # noqa: E731
BG, PANEL, LIME, TXT, MUTED, DECO, YEL = C("101827"), C("1a283b"), C("baff70"), C("f8fafc"), C("adbbcd"), C("1b2a40"), C("fbbf24")
SANS, MONO = "Arial", "Consolas"
SW, SH = Inches(13.333), Inches(7.5)
TOP = "PACKT · CLAUDE CODE BOOTCAMP COHORT 2 · LUCA BERTON"


class Deck:
    def __init__(self) -> None:
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = SW, SH
        self.blank = self.prs.slide_layouts[6]
        self.counters = []

    # ---------- primitives ----------
    def slide(self, notes: str = "", decorate: bool = True):
        s = self.prs.slides.add_slide(self.blank)
        fill = s.background.fill
        fill.solid()
        fill.fore_color.rgb = BG
        if decorate:
            self.box(s, Inches(10.9), Inches(-2.6), Inches(5.2), Inches(5.2), DECO, MSO_SHAPE.OVAL)
            self.text(s, TOP, Inches(0.7), Inches(0.35), Inches(9), Inches(0.3), 10, LIME, True)
            self.counters.append(s)
        s.notes_slide.notes_text_frame.text = notes
        return s

    def box(self, s, x, y, w, h, fill=None, shape=MSO_SHAPE.RECTANGLE, line=None):
        b = s.shapes.add_shape(shape, x, y, w, h)
        if fill is None:
            b.fill.background()
        else:
            b.fill.solid()
            b.fill.fore_color.rgb = fill
        if line is None:
            b.line.fill.background()
        else:
            b.line.color.rgb = line
            b.line.width = Pt(1.25)
        b.shadow.inherit = False
        return b

    def text(self, s, text, x, y, w, h, size, color, bold=False, align=PP_ALIGN.LEFT, font=SANS, anchor=None):
        t = s.shapes.add_textbox(x, y, w, h)
        tf = t.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.word_wrap = True
        if anchor:
            tf.vertical_anchor = anchor
        for i, line in enumerate(str(text).split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            r = p.add_run()
            r.text = line if font == MONO else line.replace("`", "")
            r.font.size, r.font.bold, r.font.name, r.font.color.rgb = Pt(size), bold, font, color
        return t

    def heading(self, s, kicker, title, y=Inches(0.95), size=34):
        self.text(s, kicker.upper(), Inches(0.7), y, Inches(11.5), Inches(0.35), 13, LIME, True)
        lines = max(1, sum(1 + len(l) // 52 for l in title.split("\n")))
        self.text(s, title, Inches(0.7), y + Inches(0.45), Inches(11.9), Inches(0.6) * lines, size, TXT, True)
        return y + Inches(0.45) + Inches(0.62) * lines + Inches(0.15)

    def panel_text(self, s, lines, x, y, w, h, size=18, mono=False, color=MUTED):
        b = self.box(s, x, y, w, h, PANEL, MSO_SHAPE.ROUNDED_RECTANGLE)
        b.adjustments[0] = 0.04
        tf = b.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.3)
        tf.margin_top = tf.margin_bottom = Inches(0.18)
        tf.vertical_anchor = MSO_ANCHOR.TOP
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.LEFT
            p.space_after = Pt(4 if mono else 8)
            r = p.add_run()
            r.text = line if mono else line.replace("`", "")
            r.font.size, r.font.name, r.font.color.rgb = Pt(size), MONO if mono else SANS, color
        return b

    # ---------- slide types ----------
    def s_title(self, d):
        s = self.slide(d["notes"])
        self.text(s, d["kicker"].upper(), Inches(0.7), Inches(1.5), Inches(11), Inches(0.4), 14, LIME, True)
        self.text(s, d["title"], Inches(0.7), Inches(2.0), Inches(11.5), Inches(2.2), 54, TXT, True)
        self.text(s, d.get("accent", ""), Inches(0.7), Inches(4.3), Inches(11.5), Inches(0.8), 36, LIME, True)
        self.text(s, d.get("sub", ""), Inches(0.7), Inches(5.4), Inches(11.5), Inches(0.5), 20, MUTED)

    def s_statement(self, d):
        s = self.slide(d["notes"])
        self.text(s, d["kicker"].upper(), Inches(0.7), Inches(1.8), Inches(11), Inches(0.4), 14, LIME, True)
        self.text(s, d["title"], Inches(0.7), Inches(2.4), Inches(11.8), Inches(2.4), 40, TXT, True)
        self.text(s, d.get("sub", ""), Inches(0.7), Inches(5.0), Inches(11.8), Inches(1.2), 20, MUTED)

    def s_section(self, d):
        s = self.slide(d["notes"])
        self.text(s, d["kicker"].upper(), Inches(0.7), Inches(2.4), Inches(11), Inches(0.4), 16, LIME, True)
        self.text(s, d["title"], Inches(0.7), Inches(2.95), Inches(11.8), Inches(1.6), 46, TXT, True)
        if d.get("sub"):
            self.text(s, d["sub"], Inches(0.7), Inches(4.7), Inches(11.8), Inches(1.0), 20, MUTED)

    def s_bullets(self, d):
        s = self.slide(d["notes"])
        y = self.heading(s, d["kicker"], d["title"])
        lines = [b if b.startswith("✓") else "•  " + b for b in d["bullets"]]
        self.panel_text(s, lines, Inches(0.7), y, Inches(11.9), SH - y - Inches(0.6), 22, color=TXT)

    def s_code(self, d):
        s = self.slide(d["notes"])
        y = self.heading(s, d["kicker"], d["title"], size=30)
        lines = d["code"].split("\n")
        longest = max(map(len, lines))
        size = 20 if len(lines) <= 10 and longest <= 62 else 17 if len(lines) <= 14 and longest <= 70 else 14
        h = SH - y - Inches(1.1 if d.get("caption") else 0.6)
        self.panel_text(s, lines, Inches(0.7), y, Inches(11.9), h, size, mono=True, color=LIME)
        if d.get("caption"):
            self.text(s, d["caption"], Inches(0.7), SH - Inches(0.9), Inches(11.9), Inches(0.4), 14, MUTED)

    def s_table(self, d):
        s = self.slide(d["notes"])
        y = self.heading(s, d["kicker"], d["title"], size=30)
        rows, cols = len(d["rows"]) + 1, len(d["header"])
        bottom = Inches(1.0) if d.get("caption") else Inches(0.5)
        h = min(SH - y - bottom, Inches(0.52) * rows)
        if d.get("caption"):
            self.text(s, d["caption"], Inches(0.7), SH - Inches(0.85), Inches(11.9), Inches(0.4), 14, MUTED)
        tbl = s.shapes.add_table(rows, cols, Inches(0.7), y, Inches(11.9), h).table
        tbl.first_row = False
        tbl.horz_banding = False
        widths = d.get("widths")
        if not widths:
            longest = [max(len(str(r[c])) for r in [d["header"], *d["rows"]]) for c in range(cols)]
            words = [max(len(w) for r in [d["header"], *d["rows"]] for w in (str(r[c]).split() or [""])) for c in range(cols)]
            raw = [max(8, l, words[c] * 1.6) for c, l in enumerate(longest)]
            widths = [max(0.12, r / sum(raw)) for r in raw]
            widths = [w / sum(widths) for w in widths]
        for c, wf in enumerate(widths):
            tbl.columns[c].width = int(Inches(11.9) * wf)
        size = 16 if rows <= 7 else 14
        for r in range(rows):
            for c in range(cols):
                cell = tbl.cell(r, c)
                cell.fill.solid()
                cell.fill.fore_color.rgb = C("22324a") if r == 0 else (PANEL if r % 2 else C("15233a"))
                cell.margin_left = cell.margin_right = Inches(0.12)
                tf = cell.text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                run = p.add_run()
                run.text = str(d["header"][c] if r == 0 else d["rows"][r - 1][c]).replace("`", "")
                run.font.size = Pt(size)
                run.font.name = SANS
                run.font.bold = r == 0 or c == 0
                run.font.color.rgb = LIME if r == 0 else (TXT if c == 0 else MUTED)

    def s_compare(self, d):
        s = self.slide(d["notes"])
        y = self.heading(s, d["kicker"], d["title"])
        h = SH - y - Inches(0.6)
        for i, (head, items) in enumerate((d["left"], d["right"])):
            x = Inches(0.7) + i * Inches(6.05)
            self.text(s, head.upper(), x, y, Inches(5.85), Inches(0.4), 15, LIME if i else YEL, True)
            self.panel_text(s, ["•  " + it for it in items], x, y + Inches(0.5), Inches(5.85), h - Inches(0.5), 19, color=TXT)

    def s_flow(self, d):
        s = self.slide(d["notes"])
        y = self.heading(s, d["kicker"], d["title"], size=30)
        n = len(d["steps"])
        w = (Inches(11.9) - Inches(0.25) * (n - 1)) / n
        for i, step in enumerate(d["steps"]):
            b = self.box(s, Inches(0.7) + i * (w + Inches(0.25)), y + Inches(0.3), int(w), Inches(1.3), PANEL, MSO_SHAPE.CHEVRON if i else MSO_SHAPE.PENTAGON)
            tf = b.text_frame
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            r = p.add_run()
            r.text = step
            r.font.size, r.font.bold, r.font.name, r.font.color.rgb = Pt(15 if n > 4 else 18), True, SANS, LIME
            tf.word_wrap = False
        self.text(s, d.get("sub", ""), Inches(0.7), y + Inches(2.1), Inches(11.9), Inches(1.4), 22, MUTED)

    def picture(self, s, path, x, y, w, h):
        from PIL import Image
        img = Image.open(path)
        iw, ih = img.size
        if ih > iw * 1.05:
            keep = min(ih, int(iw * (844 / 390 if iw < 600 else 0.72)))
            POSTERS.mkdir(exist_ok=True)
            path = POSTERS / f"crop-{Path(path).parent.name}-{Path(path).name}"
            img.crop((0, 0, iw, keep)).save(path)
            iw, ih = iw, keep
        scale = min(w / iw, h / ih)
        pw, ph = int(iw * scale), int(ih * scale)
        s.shapes.add_picture(str(path), x + (w - pw) // 2, y, pw, ph)

    def s_image(self, d):
        s = self.slide(d["notes"])
        y = self.heading(s, d["kicker"], d["title"], size=28)
        cap = Inches(0.6) if d.get("caption") else 0
        self.picture(s, ROOT / d["image"], Inches(0.7), y, Inches(11.9), SH - y - Inches(0.35) - cap)
        if d.get("caption"):
            self.text(s, d["caption"], Inches(0.7), SH - Inches(0.75), Inches(11.9), Inches(0.4), 14, MUTED)

    def s_images(self, d):
        s = self.slide(d["notes"])
        y = self.heading(s, d["kicker"], d["title"], size=28)
        n = len(d["images"])
        w = (Inches(11.9) - Inches(0.3) * (n - 1)) / n
        for i, img in enumerate(d["images"]):
            self.picture(s, ROOT / img, Inches(0.7) + i * int(w + Inches(0.3)), y, int(w), SH - y - Inches(1.0))
        self.text(s, d.get("caption", ""), Inches(0.7), SH - Inches(0.75), Inches(11.9), Inches(0.4), 14, MUTED)

    def s_video(self, part):
        video = ROOT / "videos" / f"{part['slug']}.mp4"
        notes = (f"VIDEO · Part {part['n']:02d} · {part['video']}\n"
                 "Say first: “This is a recording of a real run of the exact prompts on the previous slide. Tool calls and output are real; long runs are condensed.”\n"
                 "CLICK THE VIDEO to play. Click again to pause for questions. Then → for what the evidence shows.\n"
                 f"Fallback if playback fails: open videos/{video.name} in VLC or a browser."
                 + ("\nThis video covers prompt 1 only: Claude in Chrome cannot run headless. Do prompt 2 live with `claude --chrome`, or show the captured screenshots on the next slides." if part["n"] == 7 else ""))
        s = self.slide(notes, decorate=False)
        probe = subprocess.run(["ffprobe", "-v", "error", "-show_entries", "format=duration", "-of", "csv=p=0", str(video)],
                               capture_output=True, text=True) if video.exists() else None
        if probe is None or probe.returncode != 0 or not probe.stdout.strip():
            self.text(s, f"VIDEO PENDING: videos/{video.name}", Inches(0.7), Inches(3.3), Inches(11.9), Inches(0.8), 30, YEL, True)
            return
        POSTERS.mkdir(exist_ok=True)
        poster = POSTERS / f"{part['slug']}.png"
        subprocess.run(["ffmpeg", "-loglevel", "error", "-y", "-ss", "1.5", "-i", str(video), "-frames:v", "1", str(poster)], check=True)
        s.shapes.add_movie(str(video), 0, 0, SW, SH, poster_frame_image=str(poster), mime_type="video/mp4")

    # ---------- per-part slides ----------
    def part(self, p, start, end):
        nn = f"{p['n']:02d}"
        self.s_section({"kicker": f"Part {nn} · {start}–{end} · {p['minutes']} min", "title": p["title"],
                        "sub": p["project"],
                        "notes": f"{start}–{end}. Part {nn}: {p['title']}. {p['project']}. By the end of this part you leave with: " + "; ".join(p["outputs"]) + "."})
        self.s_bullets({"kicker": f"Part {nn} · goals", "title": "What you will be able to do", "bullets": p["goals"],
                        "notes": "These are the goals for this part. " + " ".join(p["goals"]) + "."})
        for d in CONCEPTS.get(p["n"], []):
            getattr(self, f"s_{d['type']}")(d)
        for i, (label, prompt) in enumerate(p["prompts"], start=1):
            self.s_code({"kicker": f"Part {nn} · prompt {i} of {len(p['prompts'])}", "title": label,
                         "code": wrap_prompt(prompt), "caption": f"Copy from parts/{p['slug']}/PROMPTS.md",
                         "notes": f"This is the exact prompt {i}. Read it with me: notice the scope, the constraints and the evidence it asks for. " + prompt})
        self.s_video(p)
        self.s_evidence(p)
        self.s_bullets({"kicker": f"Part {nn} · discuss", "title": "Your turn to think", "bullets": p["discussion"],
                        "notes": "Discussion. Take one of these questions. I'll take two or three answers before the lab."})
        self.s_bullets({"kicker": f"Part {nn} · lab", "title": "Hands-on",
                        "bullets": [p["lab"], f"Start: bin/bootcamp stage {nn}", "Prompts: parts/" + p["slug"] + "/PROMPTS.md",
                                    "Check: bin/bootcamp verify"],
                        "notes": "Lab time. Stage the part, run the prompts, and check with verify. If you get stuck, fallback loads the reference."})
        self.s_bullets({"kicker": f"Part {nn} · checkpoint", "title": "You now have",
                        "bullets": [f"✓  {o}" for o in p["outputs"]] + [f"Behind? bin/bootcamp checkpoint {nn} --force"],
                        "notes": "Checkpoint. If you have all of these, you're ready for the next part. If not, checkpoint gets you there in one command."})

    def s_evidence(self, p):
        nn = f"{p['n']:02d}"
        ev = ROOT / "parts" / p["slug"] / "evidence" / "recorded-run"
        right = ["(no recording yet)"]
        notes = "Now let's look at what we actually saw. Each line on the right comes from the recorded run, not from me."
        if (ev / "summary.json").exists():
            f = facts(ev)
            mins, secs = divmod(f["seconds"], 60)
            right = [f"{f['turns']} prompt(s) · {mins} min {secs:02d} s · ${f['cost']:.2f} · {f['tool_calls']} tool calls"]
            if f["tests"]:
                right.append("Tests seen in the run: " + " · ".join(f["tests"]))
            right.append(f"bin/bootcamp verify: {f['verify_pass']} PASS · {f['verify_fail']} FAIL" +
                         (f" ({'; '.join(f['fail_names'])})" if f["fail_names"] else ""))
            changed = [c.strip() for c in f["changed"]][:4]
            if changed:
                right.append("Changed: " + ", ".join(changed) + (f" … (+{len(f['changed']) - 4})" if len(f["changed"]) > 4 else ""))
            notes += f" The run took {mins} minutes {secs} seconds and cost {f['cost']:.2f} dollars."
            if f["verify_fail"]:
                notes += " Note the failing check: " + "; ".join(f["fail_names"]) + ". We discuss why rather than hide it."
        notes += " " + EXTRA_EVIDENCE_NOTES.get(p["n"], "")
        self.s_compare({"kicker": f"Part {nn} · evidence, not claims", "title": "What the run proved",
                        "left": ("What to look for", p["evidence"]), "right": ("What the recorded run produced", right),
                        "notes": notes})

    def s_break(self, label, start, end, minutes, after):
        self.s_section({"kicker": f"{start}–{end}", "title": f"{label} · {minutes} minutes",
                        "sub": f"Behind? bin/bootcamp checkpoint {after:02d} --force  ·  next: Part {after + 1:02d}",
                        "notes": BREAK_NOTES[label]})

    def build(self):
        for d in INTRO:
            getattr(self, f"s_{d['type']}")(d)
        agenda = [[s + "–" + e, (f"Part {b['n']:02d} · {b['title']}" if k == "part" else b), str(m) + " min"]
                  for k, b, s, e, m in schedule()]
        cut = next(i for i, r in enumerate(agenda) if r[1] == "Lunch") + 1
        for label, rows in (("morning", agenda[:cut]), ("afternoon", agenda[cut:])):
            self.s_table({"kicker": f"Exact schedule · {label}", "title": "Today, minute by minute", "header": ["Time", "Session", "Length"],
                          "rows": rows, "widths": [0.18, 0.67, 0.15],
                          "notes": f"This is the exact {label} schedule from the brief. The whole day is 450 minutes including breaks and lunch."})
        last = 0
        for kind, b, start, end, minutes in schedule():
            if kind == "part":
                self.part(b, start, end)
                last = b["n"]
            elif kind == "break":
                self.s_break(b, start, end, minutes, last)
        for d in FINAL:
            getattr(self, f"s_{d['type']}")(d)
        slides = list(self.prs.slides)
        total = len(slides)
        index = {id(sl): n for n, sl in enumerate(slides, start=1)}
        for s in self.counters:
            self.text(s, f"{index[id(s)]:03d} / {total}", Inches(11.4), Inches(7.05), Inches(1.4), Inches(0.3), 10, MUTED, align=PP_ALIGN.RIGHT)
        self.prs.save(OUT)
        print(f"{OUT.relative_to(ROOT)}: {len(self.prs.slides)} slides, {OUT.stat().st_size / 1e6:.1f} MB")


EXTRA_EVIDENCE_NOTES = {
    2: "In the first recording Claude refused to invent CLI output when the scratch path was blocked. Remember that: no evidence, no claim.",
    3: "Headless runs cannot approve edits under .claude/, so the rules were installed by a labelled human step. Live, you approve them.",
    5: "Watch how /goal kept the turn going until the checklist had real evidence.",
    6: "main stayed untouched: every change is on the worktree branch.",
    7: "The Chrome step can't run headless; the browser evidence was captured with the same seeded data.",
    9: "The read-only reviewer found more issues than our reference, including an unescaped status and the test files served publicly.",
    10: "The plugin's own change-review caught that the hook scripts were not executable. The one failing check is that real gap.",
    11: "Headless mode coordinated parallel background agents with plan approval and SendMessage steering. The interactive team UI is the live demo.",
    12: "Our reference says conditional go; the recorded review said no-go for the org-wide rollout. Two reviewers, two decisions: which would you sign?",
    13: "No invented deadlines: the run wrote 'not set' because the sources give none.",
}


def wrap_prompt(prompt: str, width: int = 60) -> str:
    import textwrap
    return "\n".join(textwrap.wrap(prompt, width))


if __name__ == "__main__":
    Deck().build()
